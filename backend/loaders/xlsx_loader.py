"""M1 本地数据装载与标准化（PRD §6.1 M1）。

- CSV/XLSX/PDF loader 可替换：本类实现 XLSX；`Loader` 协议约束替换位（B 轨仅换 loader，计算零改动）。
- 96 点入库；省间滚撮 24→96 展开（小时值复制 4 份）。
- 三类标注（披露/推断/补齐 + 来源日）；缺数补齐：均值/插值兜底（人工填数/相似日由修订链与 M2 承担）。
- **数据底账合并**：基线文件 + 用户上传表（backend/data/uploads/，按日期合并、后传覆盖同日）。
- **滚撮日（D 日）动态化**：运行日政策（不可披露项最近日法/核电 min 标注）在 `apply_run_day_policy`
  中按所选 D 日执行，装载期保持原始口径；候选日与 A 日解析见 `valid_target_days` / `resolve_a_day`。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

import openpyxl

from ..config import DA_SHEETS, DA_XLSX, RT_SHEETS, RT_XLSX, ROLL_AUCTION_XLSX
from ..mock.roll_auction import synth_roll_auction_24

N = 96


@dataclass
class Series:
    """一条 96/24 点序列 + 逐点标注（kind: 披露/推断/补齐/缺输入/意向/计算）。"""
    values: list[float | None]
    kinds: list[str] = field(default_factory=list)
    src_days: list[str | None] = field(default_factory=list)
    notes: list[str | None] = field(default_factory=list)

    @staticmethod
    def make(values: list[float | None], kind: str = "披露", src_day: str | None = None, note: str | None = None) -> "Series":
        n = len(values)
        return Series(values, [kind] * n, [src_day] * n, [note] * n)

    def annotate(self, i: int, kind: str, src_day: str | None = None, note: str | None = None) -> None:
        self.kinds[i] = kind
        self.src_days[i] = src_day
        self.notes[i] = note

    def to_json(self) -> dict:
        return {"values": self.values, "kinds": self.kinds, "srcDays": self.src_days, "notes": self.notes}


@dataclass
class LoadedData:
    """全量装载结果：逐日边界 + 省间滚撮 + 实时侧（原始口径，未应用运行日政策）。"""
    days: list[str]                                   # 库内全部日期（升序）
    day_ahead: dict[str, dict[str, Series]]           # day -> sheet -> Series
    realtime: dict[str, dict[str, Series]]            # day -> sheet -> Series
    roll_auction: dict[str, dict[str, list[float]]]   # day -> {volume24, price24}
    roll_source: str
    warnings: list[str] = field(default_factory=list)


class Loader(Protocol):
    """loader 可替换接口（B 轨外部通道仅替换本协议实现）。"""

    def load(self) -> LoadedData: ...


def read_sheet(path: Path, sheet: str) -> dict[str, list[float | None]]:
    """读取 32×97（或 32×25）sheet → {date: 值列表}；None 保留为缺数。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return {}
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return {}
        n_vals = len(header) - 1
        out: dict[str, list[float | None]] = {}
        for r in rows:
            if r is None or r[0] is None or not hasattr(r[0], "year"):
                continue
            key = f"{r[0].year:04d}-{r[0].month:02d}-{r[0].day:02d}"
            vals: list[float | None] = []
            for v in r[1:n_vals + 1]:
                vals.append(round(float(v), 4) if isinstance(v, (int, float)) else None)
            out[key] = vals
        return out
    finally:
        wb.close()


def fill_missing(values: list[float | None]) -> tuple[list[float | None], list[str], list[str | None], list[str | None]]:
    """缺数补齐：线性插值 → 全日均值 → 缺输入（有据标 `▣补齐(方式+依据)`，无据标"缺输入"）。"""
    n = len(values)
    out = list(values)
    kinds = ["披露" if v is not None else "" for v in values]
    src_days: list[str | None] = [None] * n
    notes: list[str | None] = [None] * n
    idx_known = [i for i, v in enumerate(values) if v is not None]
    for i, v in enumerate(values):
        if v is not None:
            continue
        prev = max((j for j in idx_known if j < i), default=None)
        nxt = min((j for j in idx_known if j > i), default=None)
        if prev is not None and nxt is not None:
            w = (i - prev) / (nxt - prev)
            out[i] = round(values[prev] + w * (values[nxt] - values[prev]), 4)   # type: ignore[index]
            kinds[i] = "补齐"
            notes[i] = "▣补齐(线性插值)"
        elif idx_known:
            avg = sum(values[j] for j in idx_known) / len(idx_known)              # type: ignore[misc]
            out[i] = round(avg, 4)
            kinds[i] = "补齐"
            notes[i] = "▣补齐(全日均值)"
        else:
            kinds[i] = "缺输入"
            notes[i] = "无据可补"
    return out, kinds, src_days, notes


def _merge_sheets(base: dict[str, dict[str, list[float | None]]], extra: dict[str, dict[str, list[float | None]]]) -> None:
    """按日期合并：extra 中同 (sheet, day) 覆盖 base（后传覆盖同日）。"""
    for sheet, by_day in extra.items():
        for day, vals in by_day.items():
            base.setdefault(sheet, {})[day] = vals


class XlsxLoader:
    """A 轨实现：基线 xlsx + 上传表 → LoadedData（M1 原始口径）。"""

    def __init__(self, da_path: Path = DA_XLSX, rt_path: Path = RT_XLSX, roll_path: Path | None = None,
                 extra_da: tuple[Path, ...] = (), extra_rt: tuple[Path, ...] = (),
                 extra_roll: tuple[Path, ...] = ()) -> None:
        self.da_path = da_path
        self.rt_path = rt_path
        self.roll_path = roll_path if roll_path is not None else ROLL_AUCTION_XLSX
        self.extra_da = tuple(extra_da)
        self.extra_rt = tuple(extra_rt)
        self.extra_roll = tuple(extra_roll)

    def load(self) -> LoadedData:
        warnings: list[str] = []
        raw_da: dict[str, dict[str, list[float | None]]] = {s: read_sheet(self.da_path, s) for s in DA_SHEETS}
        raw_rt: dict[str, dict[str, list[float | None]]] = {s: read_sheet(self.rt_path, s) for s in RT_SHEETS}

        # ---- 用户上传合并（时间升序 = 后传覆盖同日）----
        for f in self.extra_da:
            for s in DA_SHEETS:
                _merge_sheets(raw_da, {s: read_sheet(f, s)})
        for f in self.extra_rt:
            for s in RT_SHEETS:
                _merge_sheets(raw_rt, {s: read_sheet(f, s)})
        if self.extra_da or self.extra_rt:
            warnings.append(f"已合并上传表：日前×{len(self.extra_da)}、实时×{len(self.extra_rt)}（按日期覆盖）")

        all_days = sorted(set(raw_da.get("负荷", {}).keys()) | set(raw_rt.get("实时电价", {}).keys()))
        if not all_days:
            raise RuntimeError(f"数据底账为空：基线 {self.da_path} 无负荷数据且无有效上传")
        days = all_days

        # ---- 日前/实时逐日入库（补齐 + 披露标注）----
        day_ahead: dict[str, dict[str, Series]] = {}
        for day in days:
            day_ahead[day] = {}
            for sheet, by_day in raw_da.items():
                if not by_day:
                    continue
                if day not in by_day or not by_day[day]:
                    continue                     # 该日缺此 sheet → 不造缺输入占位（候选日判定按需检查）
                vals_raw = by_day[day]
                if len(vals_raw) == 24:          # 24 点原生 sheet 保持 24 点
                    vals = (vals_raw + [None] * 24)[:24]
                else:
                    vals = (vals_raw + [None] * N)[:N]
                vals, kinds, srcs, notes = fill_missing(vals)
                day_ahead[day][sheet] = Series(vals, kinds, srcs, notes)

        realtime: dict[str, dict[str, Series]] = {}
        for sheet, by_day in raw_rt.items():
            if not by_day:
                continue
            for day, vals_raw in by_day.items():
                vals = (vals_raw + [None] * N)[:N]
                vals, kinds, srcs, notes = fill_missing(vals)
                realtime.setdefault(day, {})[sheet] = Series(vals, kinds, srcs, notes)

        # ---- 省间滚撮：真实文件/上传优先（A-0 替换位），否则确定性合成占位 ----
        roll: dict[str, dict[str, list[float]]] = {}
        roll_files = [self.roll_path] if self.roll_path.exists() else []
        roll_files += list(self.extra_roll)
        if roll_files:
            roll_source = "、".join(str(p) for p in roll_files)
            for f in roll_files:
                vol_by_day = read_sheet(f, "成交量")
                price_by_day = read_sheet(f, "价格")
                for day in days:
                    v24 = (vol_by_day.get(day, []) + [0.0] * 24)[:24]
                    p24 = (price_by_day.get(day, []) + [0.0] * 24)[:24]
                    if day in vol_by_day:
                        roll[day] = {"volume24": v24, "price24": p24}
        else:
            roll_source = "确定性合成占位（省间滚撮 24 点本地暂缺，A-0 待业务方提供；可在数据管理上传滚撮表）"
            for day in days:
                tie = day_ahead.get(day, {}).get("联络线")
                tie_vals = tie.values if tie is not None else [None] * N
                v24, p24 = synth_roll_auction_24(tie_vals)
                roll[day] = {"volume24": v24, "price24": p24}
            warnings.append(roll_source)

        return LoadedData(days=days, day_ahead=day_ahead, realtime=realtime,
                          roll_auction=roll, roll_source=roll_source, warnings=warnings)


# ======================= 滚撮日（D 日）动态化 =======================

# 运行日不可披露项 → 最近日法（PRD D-5；按所选 D 日以 A 日（最近披露日）替代）
UNDISCLOSED_KEYS = ("非市场化", "水电", "地方燃煤")


def valid_target_days(data: LoadedData) -> list[str]:
    """可选滚撮日 = 库内 8 项边界齐全（非全缺）的日期（升序）。"""
    required = ("负荷", "水电", "核电", "地方燃煤", "风电", "光伏", "联络线", "非市场化")
    out = []
    for day in data.days:
        da = data.day_ahead.get(day, {})
        ok = True
        for k in required:
            s = da.get(k)
            if s is None or not any(v is not None for v in s.values):
                ok = False
                break
        if ok:
            out.append(day)
    return out


def resolve_a_day(data: LoadedData, d_day: str) -> str:
    """A 日 = D 日之前、库内有日前电价的最新日（"数据库中该日之前的最新数据"）。"""
    candidates = [d for d in data.days
                  if d < d_day and any(v is not None for v in data.day_ahead.get(d, {}).get("日前电价", Series([])).values)]
    if not candidates:
        raise ValueError(f"{d_day} 之前库内无日前电价日，无法拟合（A 日缺失）")
    return max(candidates)


def apply_run_day_policy(data: LoadedData, d_day: str, a_day: str) -> list[str]:
    """运行日政策（在所选 D 日副本上执行；返回 warnings）：
    1) 不可披露项（非市场化/水电/地方燃煤）最近日法 → 取 A 日披露 + 标注来源日；
    2) 核电预测 = min(最新披露运行日出力, 检修上限=∞（检修计划缺，未检修默认开机））→ 标注。"""
    warnings: list[str] = []
    da_d = data.day_ahead.get(d_day, {})
    da_a = data.day_ahead.get(a_day, {})
    for key in UNDISCLOSED_KEYS:
        tgt, src = da_d.get(key), da_a.get(key)
        if tgt is None or src is None:
            warnings.append(f"{d_day} 缺边界「{key}」或 {a_day} 缺披露源，最近日法未应用")
            continue
        for i in range(min(len(tgt.values), len(src.values))):
            if src.values[i] is not None:
                tgt.values[i] = src.values[i]
                tgt.annotate(i, "推断", a_day, f"运行日不披露，最近日法取自 {a_day}")
    nuc = da_d.get("核电")
    if nuc is not None:
        for i in range(len(nuc.values)):
            if nuc.values[i] is not None:
                nuc.annotate(i, "推断", d_day, "核电预测 = min(最新披露运行日出力, 检修上限=∞（检修缺，未检修默认开机）)")
    return warnings


def expand24to96(v24: list[float | None]) -> list[float | None]:
    """24→96 展开：小时值复制 4 份（业务方确认按复制实现）。"""
    out: list[float | None] = []
    for v in v24:
        out.extend([v] * 4)
    return out


def avg96to24(v96: list[float | None]) -> list[float | None]:
    """96→24 合成：每 4 点等权算术平均（缺数跳过；组内全缺 → None）。"""
    out: list[float | None] = []
    for i in range(24):
        seg = [x for x in v96[i * 4:i * 4 + 4] if x is not None]
        out.append(round(sum(seg) / len(seg), 6) if seg else None)
    return out
