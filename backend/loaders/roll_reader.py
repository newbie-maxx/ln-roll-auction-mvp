"""省间滚撮表读取适配器：兼容两种真实布局。

- 长表（本次业务方文件 省间滚撮统计8月.xlsx）：单 sheet「成交量」，行 =（日期+小时, 成交量），
  每日 24 行；小时 h 视为 h:00–h+1:00 时段（period = h+1，实现口径待业务确认）。「价格」sheet 可选。
- 宽表（上传契约）：sheets「成交量」/「价格」，行 = 日期，其后 24 点数值列。

价格 sheet 缺失 → volume 用真实值、price 返回 None（由 loader 降级为合成占位 + 警示；价格仅用于
快照展示，不进计算链——M4 基线/M5 省间交易总数只用成交量）。
"""
from __future__ import annotations

from pathlib import Path

import openpyxl


def _is_long_layout(rows: list[tuple]) -> bool:
    """首数据列含带小时的 datetime → 长表（行=日期+小时）。"""
    for r in rows[:50]:
        v = r[0] if r else None
        if v is not None and hasattr(v, "hour") and getattr(v, "hour", 0) != 0:
            return True
        if v is not None and hasattr(v, "hour") and getattr(v, "hour", 0) == 0:
            # 0 点行也可能出现：看是否同一天出现多行（长表特征）
            days = [x[0] for x in rows[:50] if x and hasattr(x[0], "hour")]
            if len(days) >= 3 and len({(d.year, d.month, d.day) for d in days}) <= len(days) // 2:
                return True
    return False


def _read_long(ws) -> dict[str, dict[str, list[float] | None]]:
    """长表 → {day: {"volume24": [...], "price24": None}}；缺行小时置 None（不产猜测值）。"""
    volume: dict[str, list[float | None]] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None or not hasattr(r[0], "hour"):
            continue
        d = r[0]
        day = f"{d.year:04d}-{d.month:02d}-{d.day:02d}"
        arr = volume.setdefault(day, [None] * 24)
        h = min(max(d.hour, 0), 23)
        if isinstance(r[1], (int, float)):
            arr[h] = round(float(r[1]), 4)
    return {day: {"volume24": arr, "price24": None} for day, arr in volume.items()}


def _read_wide(ws) -> dict[str, list[float]]:
    """宽表 → {day: [24 值]}（列不足 24 补 0；行=日期）。"""
    out: dict[str, list[float]] = {}
    rows = ws.iter_rows(values_only=True)
    next(rows, None)                                    # 表头
    for r in rows:
        if not r or r[0] is None or not hasattr(r[0], "year"):
            continue
        d = r[0]
        day = f"{d.year:04d}-{d.month:02d}-{d.day:02d}"
        vals = [round(float(v), 4) if isinstance(v, (int, float)) else 0.0 for v in r[1:25]]
        vals += [0.0] * (24 - len(vals))
        out[day] = vals
    return out


def read_roll_file(path: Path) -> tuple[dict[str, dict[str, list[float] | None]], str]:
    """读滚撮表（自动识别长/宽布局）；返回 ({day: {volume24, price24}}, 布局说明)。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "成交量" not in wb.sheetnames:
            return {}, "缺「成交量」sheet"
        ws = wb["成交量"]
        rows = [(r[0], r[1] if len(r) > 1 else None) for r in ws.iter_rows(values_only=True) if r]
        data: dict[str, dict[str, list[float] | None]]
        if _is_long_layout(rows):
            data = _read_long(ws)
            layout = "长表（行=日期+小时，每日 24 行；小时按 h:00–h+1:00 计时段）"
        else:
            wide = _read_wide(ws)
            data = {day: {"volume24": arr, "price24": None} for day, arr in wide.items()}
            layout = "宽表（行=日期，列=24 点）"
        # 价格 sheet（可选，仅展示用途）
        if "价格" in wb.sheetnames:
            price = _read_wide(wb["价格"])
            for day, p24 in price.items():
                data.setdefault(day, {"volume24": [None] * 24, "price24": None})["price24"] = p24
        return data, layout
    finally:
        wb.close()


def looks_like_roll_file(path: Path) -> bool:
    """sheet 签名探测：含「成交量」即视为滚撮表（日前/实时边界表无此 sheet）。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:                                    # noqa: BLE001
        return False
    try:
        return "成交量" in wb.sheetnames
    finally:
        wb.close()
